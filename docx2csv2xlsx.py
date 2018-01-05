# -*- coding: utf-8 -*-
"""
Created on Sat Oct 28 11:43:21 2017

@author: Rômulo Ponciano 
@email: github.com/rponciano
"""

import docx
import csv
from xlsxwriter.workbook import Workbook
import docx2txt

# docx para csv
def docx2csv(filename):
    # docx to list
    doc = docx.Document(filename)
    fulltext = []
    for para in doc.paragraphs:
        fulltext.append(para.text)
    # list to csv
    f = open('emails.csv', 'w')
    for item in fulltext:
        f.write(item + '\n')
    f.close()
    
# docx para txt
def doc2txt(filename):
    text = docx2txt.process(filename)
    f = open('novo.txt', 'w')
    f.write(text)
    f.close()

# csv para xlsx
def csv2xlsx(csvfile):
    workbook = Workbook(csvfile + '.xlsx')
    worksheet = workbook.add_worksheet()
    with open(csvfile, 'rt') as f:
        reader = csv.reader(f)
        for r, row in enumerate(reader):
            for c, col in enumerate(row):
                worksheet.write(r, c, col)
    workbook.close()

# if you need clean or process the csv of any way,
# do it here.    
def processar(filename):
    with open(filename) as fvelho:
        with open('emailsfinal.csv', 'w') as fnovo:
            for line in fvelho:
                lista = line.split()
                if len(lista)>1:
                    email = lista[len(lista)-1]
                    lista = lista[:-1]
                    nome = " ".join(lista)
                    fnovo.write(nome + ',' + email + '\n')
    fvelho.close()
    fnovo.close()
    
# ------------------------------------------------------------
if __name__ == '__main__':
    doc2txt('MaisSNO.docx')
    #docx2csv('MaisSNO.docx') # name of the docx
    processar('novo.txt')
    csv2xlsx('emailsfinal.csv') # name of the processed csv
