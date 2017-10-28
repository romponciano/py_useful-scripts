# -*- coding: utf-8 -*-
"""
Created on Sat Oct 28 11:43:21 2017
Cliente: Luis Gustavo
Criado para converter word docx emails e nomes
para excel. 

@author: Rômulo Ponciano 
@email: romuloponciano@id.uff.br
"""

import docx
import csv
from xlsxwriter.workbook import Workbook

# just change the name emails.csv, if you 
# want a different name.
def docx2csv(filename):
    # docx to list
    doc = docx.Document(filename)
    fulltext = []
    for para in doc.paragraphs:
        fulltext.append(para.text)
    # list to csv
    f = open('emails.csv', 'w')
    for item in fulltext:
        f.write(item + '\n')
    f.close()

# if you need clean or process the csv of any way,
# do it here.    
def processarCsv(filename):
    with open(filename) as fvelho:
        with open('emailsfinal.csv', 'w') as fnovo:
            for line in fvelho:
                linha = line.replace(' | ', ',')
                linha = linha.replace('.COM ', '.COM\n')
                linha = linha.replace('.BR ', '.BR\n')
                fnovo.write(linha)
    fvelho.close()
    fnovo.close()
    
# csv to xlsx. Don't need to change anything.
def csv2xlsx(csvfile):
    workbook = Workbook(csvfile + '.xlsx')
    worksheet = workbook.add_worksheet()
    with open(csvfile, 'rt') as f:
        reader = csv.reader(f)
        for r, row in enumerate(reader):
            for c, col in enumerate(row):
                worksheet.write(r, c, col)
    workbook.close()
    
# ------------------------------------------------------------
if __name__ == '__main__':
    docx2csv('FAZENDO.docx') # name of the docx
    processarCsv('emails.csv') # name of the first csv file
    csv2xlsx('emailsfinal.csv') # name of the processed csv